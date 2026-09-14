#!/usr/bin/env python3
"""
Rebuild "Elements of Thermal Preference" (v18) from figure_spec.json.

    pip install python-pptx
    python build_thermal_preference_figure.py

Writes thermal_preference_elements.pptx next to this script.

The spec holds every shape in drawing order: 41 freeform outlines (the occupant,
window, mug, food bowl, the four category icons and the scene line art), 27
straight connectors (leader lines, the header rules, the panel divider), 7 preset
shapes and 54 text boxes. All coordinates are in figure pixels, 1 px = 9525 EMU,
on a 1900 x 880 px canvas. At 190 mm reproduction width that puts element labels
and citations at 6.8 pt.

To change the figure, edit figure_spec.json (it is plain JSON) or patch the spec
dict in memory before build() is called -- see the tweak_example() helper.
"""
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

PX = 9525
HERE = os.path.dirname(os.path.abspath(__file__))
SPEC = os.path.join(HERE, "figure_spec.json")
OUT = os.path.join(HERE, "thermal_preference_elements.pptx")

PRESET = {"rect": MSO_SHAPE.RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
          "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE, "line": None}
ALIGN = {"LEFT": PP_ALIGN.LEFT, "RIGHT": PP_ALIGN.RIGHT, "CENTER": PP_ALIGN.CENTER}
ANCHOR = {"TOP": MSO_ANCHOR.TOP, "MIDDLE": MSO_ANCHOR.MIDDLE, "BOTTOM": MSO_ANCHOR.BOTTOM}


def E(v):
    return Emu(int(round(v * PX)))


def _style(shape, rec):
    """Apply fill and line, and strip the theme style that adds a drop shadow."""
    shape.shadow.inherit = False
    fill = rec.get("fill")
    if hasattr(shape, "fill"):                 # connectors have no fill
        if fill in (None, "none"):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    ln = rec.get("line")
    if not ln or ln.get("color") in (None, "none"):
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(ln["color"])
        if ln.get("w"):
            shape.line.width = Pt(ln["w"])
        if ln.get("dash"):
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    for end in ("headEnd", "tailEnd"):
        if ln and ln.get(end):
            e = ln[end]
            lnEl = shape.line._get_or_add_ln()
            lnEl.append(lnEl.makeelement(qn("a:" + end),
                                         {k: v for k, v in e.items() if v}))
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    spPr = shape._element.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None:
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def xfrm_attr(rec):
    """Rotation / flips recorded from the source shape."""
    a = ""
    if rec.get("rot"):
        a += ' rot="%d"' % rec["rot"]
    if rec.get("flipH"):
        a += ' flipH="1"'
    if rec.get("flipV"):
        a += ' flipV="1"'
    return a


def add_freeform(slide, rec, uid):
    """One <a:path> holding every subpath, so opposite windings cut holes."""
    paths_xml = []
    for path in rec["geom"]:
        seg = []
        flags = path.get("closed") or [True] * len(path["subs"])
        for sub, shut in zip(path["subs"], flags):
            for i, (px, py) in enumerate(sub):
                tag = "moveTo" if i == 0 else "lnTo"
                seg.append('<a:%s><a:pt x="%d" y="%d"/></a:%s>' % (tag, px, py, tag))
            if shut:                       # open strokes must not be closed
                seg.append("<a:close/>")
        paths_xml.append('<a:path w="%d" h="%d">%s</a:path>'
                         % (path["w"], path["h"], "".join(seg)))
    fill = rec.get("fill")
    fill_xml = ('<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill
                if fill and fill != "none" else "<a:noFill/>")
    ln = rec.get("line") or {}
    ends = ""
    for end in ("headEnd", "tailEnd"):          # the wavy radiation arrows carry these
        if ln.get(end):
            attrs = " ".join('%s="%s"' % (k, v) for k, v in ln[end].items() if v)
            ends += "<a:%s %s/>" % (end, attrs)
    if ln.get("color") and ln["color"] != "none":
        line_xml = ('<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>%s</a:ln>'
                    % (int(round((ln.get("w") or 1) * 12700)), ln["color"], ends))
    else:
        line_xml = "<a:ln><a:noFill/></a:ln>"
    xml = (
        '<p:sp %s>'
        '<p:nvSpPr><p:cNvPr id="%d" name="Freeform %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
        '<a:rect l="l" t="t" r="r" b="b"/><a:pathLst>%s</a:pathLst></a:custGeom>'
        '%s%s<a:effectLst/></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
    ) % (nsdecls("p", "a"), uid, uid, xfrm_attr(rec), E(rec["x"]), E(rec["y"]),
         E(rec["w"]), E(rec["h"]), "".join(paths_xml), fill_xml, line_xml)
    slide.shapes._spTree.append(parse_xml(xml))


def add_connector(slide, rec):
    x1 = rec["x"] + (rec["w"] if rec.get("flipH") else 0)
    y1 = rec["y"] + (rec["h"] if rec.get("flipV") else 0)
    x2 = rec["x"] + (0 if rec.get("flipH") else rec["w"])
    y2 = rec["y"] + (0 if rec.get("flipV") else rec["h"])
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _style(c, rec)
    return c


def add_preset(slide, rec):
    s = slide.shapes.add_shape(PRESET[rec["prst"]], E(rec["x"]), E(rec["y"]),
                               E(rec["w"]), E(rec["h"]))
    if rec.get("rot"):
        s.rotation = rec["rot"] / 60000.0
    _style(s, rec)
    return s


def add_text(slide, rec):
    tb = slide.shapes.add_textbox(E(rec["x"]), E(rec["y"]), E(rec["w"]), E(rec["h"]))
    if rec.get("rot"):
        tb.rotation = rec["rot"] / 60000.0
    tf = tb.text_frame
    tf.word_wrap = rec.get("wrap", True)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if rec.get("anchor") in ANCHOR:
        tf.vertical_anchor = ANCHOR[rec["anchor"]]
    for i, para in enumerate(rec["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if para.get("align") in ALIGN:
            p.alignment = ALIGN[para["align"]]
        for run in para["runs"]:
            r = p.add_run()
            r.text = run["t"]
            if run.get("sz"):
                r.font.size = Pt(run["sz"])
            r.font.bold = run.get("b", False)
            if run.get("font"):
                r.font.name = run["font"]
            r.font.color.rgb = RGBColor(0, 0, 0)
            if run.get("sup"):
                r.font._rPr.set("baseline", run["sup"])
    return tb


def build(spec, out=OUT):
    prs = Presentation()
    prs.slide_width = E(spec["width"])
    prs.slide_height = E(spec["height"])
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank
    uid = 1000
    for rec in spec["shapes"]:                            # drawing order = z-order
        kind = rec["kind"]
        if kind == "freeform":
            uid += 1
            add_freeform(slide, rec, uid)
        elif kind == "connector":
            add_connector(slide, rec)
        elif kind == "preset":
            add_preset(slide, rec)
        elif kind == "text":
            add_text(slide, rec)
    prs.save(out)
    return out


# --------------------------------------------------------------------------
def find(spec, text):
    """Every shape whose first line starts with `text` -- handy for edits."""
    hits = []
    for s in spec["shapes"]:
        if s["kind"] == "text" and s["paras"][0]["runs"]:
            if s["paras"][0]["runs"][0]["t"].startswith(text):
                hits.append(s)
    return hits


def tweak_example(spec):
    """Illustration only: set every citation to 20 pt and move a label down."""
    for s in spec["shapes"]:
        if s["kind"] != "text":
            continue
        for para in s["paras"]:
            for run in para["runs"]:
                if run.get("sup"):
                    run["sz"] = 20
    for s in find(spec, "Posture"):
        s["y"] += 4
    return spec


if __name__ == "__main__":
    spec = json.load(open(SPEC))
    if "--tweak" in sys.argv:
        spec = tweak_example(spec)
    path = build(spec)
    n = len(spec["shapes"])
    print(f"{n} shapes -> {path}")
    print(f"canvas {spec['width']} x {spec['height']} px "
          f"({spec['width']*0.75:.0f} x {spec['height']*0.75:.0f} pt)")
